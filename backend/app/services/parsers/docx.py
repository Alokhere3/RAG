"""DOCX document parser using python-docx."""

import io
from typing import List
import docx
from docx.opc.exceptions import PackageNotFoundError

from app.core.config import get_settings
from app.core.errors import CorruptFileError, ParserError
from app.services.parsers.base import BaseParser, PageContent, ParseResult


class DocxParser(BaseParser):
    """Parser for Microsoft Word (.docx) documents."""

    @property
    def name(self) -> str:
        return "docx-parser"

    @property
    def version(self) -> str:
        return docx.__version__

    @property
    def supported_mime_types(self) -> List[str]:
        return ["application/vnd.openxmlformats-officedocument.wordprocessingml.document"]

    def parse(self, content_bytes: bytes, filename: str) -> ParseResult:
        settings = get_settings()

        if not content_bytes:
            raise CorruptFileError(f"Cannot parse empty DOCX document: {filename}")

        try:
            stream = io.BytesIO(content_bytes)
            doc = docx.Document(stream)
        except (PackageNotFoundError, Exception) as exc:
            raise CorruptFileError(f"Corrupt or unreadable DOCX document ({filename}): {str(exc)}") from exc

        paragraph_texts: List[str] = []

        # Extract text from standard paragraphs
        for p in doc.paragraphs:
            text = p.text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "").strip()
            if text:
                paragraph_texts.append(text)

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [
                    cell.text.replace("\r\n", " ").replace("\r", " ").replace("\x00", "").strip()
                    for cell in row.cells
                ]
                row_line = " | ".join(c for c in row_cells if c)
                if row_line:
                    paragraph_texts.append(row_line)

        combined_text = "\n\n".join(paragraph_texts)

        if not combined_text.strip():
            raise CorruptFileError(f"DOCX document '{filename}' contains no extractable text content.")

        if len(combined_text) > settings.parser_max_chars:
            raise ParserError(
                f"Extracted DOCX text exceeds maximum character bound of {settings.parser_max_chars} characters.",
                code="PARSER_LIMIT_EXCEEDED",
            )

        # DOCX doesn't have fixed pagination, treat as single logical document
        pages = [
            PageContent(
                page_number=1,
                text=combined_text,
                char_count=len(combined_text),
            )
        ]

        safe_metadata = {
            "filename": filename,
            "format": "docx",
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
        }

        # Safe core properties
        try:
            core_props = doc.core_properties
            if core_props.title:
                safe_metadata["title"] = str(core_props.title)[:100]
            if core_props.author:
                safe_metadata["author"] = str(core_props.author)[:100]
        except Exception:
            pass

        return ParseResult(
            text=combined_text,
            page_count=1,
            char_count=len(combined_text),
            pages=pages,
            parser_name=self.name,
            parser_version=self.version,
            metadata=safe_metadata,
        )
